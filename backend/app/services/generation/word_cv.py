"""
Generateur Word : CV methode STAR.

Genere un CV professionnel adapte a la culture cible (locale).
Utilise la methode STAR pour les experiences :
  - Situation / Tache (contexte)
  - Actions menees
  - Resultats quantifiables
"""

from typing import Any, Dict

from docx import Document
from docx.shared import Pt

from app.core.logging import get_logger
from app.services.generation.base import DocumentGenerator
from app.services.generation.style import (
    ALIGN_CENTRE,
    ALIGN_JUSTIFIE,
    COULEUR_PRIMAIRE,
    COULEUR_SECONDAIRE,
    ajouter_espace,
    ajouter_paragraphe,
    ajouter_puce,
    ajouter_titre_section,
    configurer_marges,
    configurer_style_normal,
)

logger = get_logger("generation.word_cv")

# Traductions des titres CV par langue
CV_TITLES = {
    "fr": {
        "competences": "Competences cles",
        "experiences": "Experiences professionnelles",
        "formation": "Formation",
        "projets": "Projets",
        "informations": "Informations complementaires",
        "technical_label": "Techniques : ",
        "soft_label": "Humaines : ",
        "languages_label": "Langues : ",
        "interests_label": "Centres d interet : ",
    },
    "en": {
        "competences": "Key Skills",
        "experiences": "Professional Experience",
        "formation": "Education",
        "projets": "Projects",
        "informations": "Additional Information",
        "technical_label": "Technical: ",
        "soft_label": "Soft Skills: ",
        "languages_label": "Languages: ",
        "interests_label": "Interests: ",
    },
    "es": {
        "competences": "Competencias clave",
        "experiences": "Experiencia profesional",
        "formation": "Formacion",
        "projets": "Proyectos",
        "informations": "Informacion adicional",
        "technical_label": "Tecnicas: ",
        "soft_label": "Habilidades: ",
        "languages_label": "Idiomas: ",
        "interests_label": "Intereses: ",
    },
    "de": {
        "competences": "Schluesselqualifikationen",
        "experiences": "Berufserfahrung",
        "formation": "Ausbildung",
        "projets": "Projekte",
        "informations": "Zusaetzliche Informationen",
        "technical_label": "Technisch: ",
        "soft_label": "Sozial: ",
        "languages_label": "Sprachen: ",
        "interests_label": "Interessen: ",
    },
}


def _get_cv_titles(locale: dict) -> dict:
    """Retourne les titres CV selon la langue (fallback anglais)."""
    language_code = "en"
    if locale:
        for key in ("language_code", "language", "code"):
        
            if key in locale and locale[key]:
                language_code = str(locale[key]).lower()
                break
    return CV_TITLES.get(language_code, CV_TITLES["en"])



class WordCVGenerator(DocumentGenerator):
    name = "word_cv"

    def build(self, data: Dict[str, Any], locale: Dict[str, Any]) -> Document:
        titles = _get_cv_titles(locale)
        doc = Document()
        configurer_marges(doc)
        configurer_style_normal(doc)

        # ============================================================
        # EN-TETE
        # ============================================================
        coordonnees = data.get("coordonnees", {})
        nom_complet = coordonnees.get("nom_complet", "Candidat")

        # Nom
        p_nom = doc.add_paragraph()
        r_nom = p_nom.add_run(nom_complet.upper())
        r_nom.bold = True
        r_nom.font.size = Pt(18)
        r_nom.font.color.rgb = COULEUR_PRIMAIRE
        p_nom.alignment = ALIGN_CENTRE
        p_nom.paragraph_format.space_after = Pt(2)

        # Coordonnees
        contact_parts = []
        if coordonnees.get("telephone"):
            contact_parts.append(coordonnees["telephone"])
        if coordonnees.get("email"):
            contact_parts.append(coordonnees["email"])
        if coordonnees.get("adresse"):
            contact_parts.append(coordonnees["adresse"])

        if contact_parts:
            ajouter_paragraphe(
                doc,
                " | ".join(contact_parts),
                taille=9,
                couleur=COULEUR_SECONDAIRE,
                alignement=ALIGN_CENTRE,
                espace_apres=4,
            )

        # Liens (LinkedIn, GitHub, etc.)
        liens = coordonnees.get("liens", [])
        if liens:
            ajouter_paragraphe(
                doc,
                " | ".join(liens),
                taille=9,
                couleur=COULEUR_SECONDAIRE,
                alignement=ALIGN_CENTRE,
                espace_apres=6,
            )

        # ============================================================
        # TITRE PROFESSIONNEL + ACCROCHE
        # ============================================================
        titre = data.get("titre_professionnel", "")
        if titre:
            ajouter_paragraphe(
                doc,
                titre.upper(),
                taille=12,
                gras=True,
                couleur=COULEUR_PRIMAIRE,
                alignement=ALIGN_CENTRE,
                espace_apres=4,
            )

        accroche = data.get("accroche", "")
        if accroche:
            ajouter_paragraphe(
                doc,
                accroche,
                taille=10,
                italique=True,
                alignement=ALIGN_JUSTIFIE,
                espace_apres=8,
            )

        # ============================================================
        # COMPETENCES CLES
        # ============================================================
        tech_skills = data.get("competences_techniques", [])
        hum_skills = data.get("competences_humaines", [])

        if tech_skills or hum_skills:
            ajouter_titre_section(doc, titles["competences"])

            if tech_skills:
                p = doc.add_paragraph()
                r = p.add_run(titles["technical_label"])
                r.bold = True
                r.font.size = Pt(10)
                p.add_run(", ".join(tech_skills))
                p.paragraph_format.space_after = Pt(2)

            if hum_skills:
                p = doc.add_paragraph()
                r = p.add_run(titles["soft_label"])
                r.bold = True
                r.font.size = Pt(10)
                p.add_run(", ".join(hum_skills))
                p.paragraph_format.space_after = Pt(2)

        # ============================================================
        # EXPERIENCES (METHODE STAR)
        # ============================================================
        experiences = data.get("experiences", [])
        if experiences:
            ajouter_titre_section(doc, titles["experiences"])

            for exp in experiences:
                # Ligne 1 : Poste + Entreprise
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)

                poste = exp.get("poste", "")
                entreprise = exp.get("entreprise", "")
                lieu = exp.get("lieu", "")
                periode = exp.get("periode", "")

                r_poste = p.add_run(poste)
                r_poste.bold = True
                r_poste.font.size = Pt(11)
                r_poste.font.color.rgb = COULEUR_PRIMAIRE

                if entreprise:
                    r_sep = p.add_run(" — ")
                    r_sep.font.size = Pt(10)
                    r_ent = p.add_run(entreprise)
                    r_ent.italic = True
                    r_ent.font.size = Pt(10)

                if lieu or periode:
                    details_parts = [x for x in [lieu, periode] if x]
                    r_details = p.add_run(" (" + ", ".join(details_parts) + ")")
                    r_details.italic = True
                    r_details.font.size = Pt(9)
                    r_details.font.color.rgb = COULEUR_SECONDAIRE

                # Contexte (Situation / Tache)
                situation = exp.get("situation_tache", "")
                if situation:
                    p_ctx = doc.add_paragraph()
                    p_ctx.paragraph_format.left_indent = Pt(10)
                    p_ctx.paragraph_format.space_after = Pt(2)
                    r_label = p_ctx.add_run("[Contexte] ")
                    r_label.bold = True
                    r_label.font.size = Pt(9)
                    r_ctx = p_ctx.add_run(situation)
                    r_ctx.font.size = Pt(9)

                # Actions menees (puces)
                actions = exp.get("actions_menees", [])
                for action in actions:
                    p_action = doc.add_paragraph(style="List Bullet")
                    p_action.paragraph_format.left_indent = Pt(20)
                    p_action.paragraph_format.space_after = Pt(1)
                    r_action = p_action.add_run(action)
                    r_action.font.size = Pt(9)

                # Resultats chiffres
                resultats = exp.get("resultats_quantifiables", "")
                if resultats:
                    p_res = doc.add_paragraph()
                    p_res.paragraph_format.left_indent = Pt(10)
                    p_res.paragraph_format.space_after = Pt(4)
                    r_label = p_res.add_run("[Impact] ")
                    r_label.bold = True
                    r_label.font.size = Pt(9)
                    r_label.font.color.rgb = COULEUR_PRIMAIRE
                    r_res = p_res.add_run(resultats)
                    r_res.font.size = Pt(9)

        # ============================================================
        # FORMATIONS
        # ============================================================
        formations = data.get("formations", [])
        if formations:
            ajouter_titre_section(doc, titles["formation"])

            for f in formations:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Pt(10)
                p.paragraph_format.space_after = Pt(1)

                diplome = f.get("diplome", "")
                etablissement = f.get("etablissement", "")
                annee = f.get("annee", "")

                r_dip = p.add_run(diplome)
                r_dip.bold = True
                r_dip.font.size = Pt(9)

                if etablissement:
                    r_et = p.add_run(f" — {etablissement}")
                    r_et.font.size = Pt(9)

                if annee:
                    r_an = p.add_run(f" ({annee})")
                    r_an.italic = True
                    r_an.font.size = Pt(9)

        # ============================================================
        # PROJETS
        # ============================================================
        projets = data.get("projets", [])
        if projets:
            ajouter_titre_section(doc, titles["projets"])

            for proj in projets:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Pt(10)
                p.paragraph_format.space_after = Pt(1)

                nom = proj.get("nom", "")
                description = proj.get("description", "")
                technos = proj.get("technologies", [])

                r_nom = p.add_run(nom)
                r_nom.bold = True
                r_nom.font.size = Pt(9)

                if description:
                    r_desc = p.add_run(f" : {description}")
                    r_desc.font.size = Pt(9)

                if technos:
                    r_tech = p.add_run(f" [{", ".join(technos)}]")
                    r_tech.italic = True
                    r_tech.font.size = Pt(8)

        # ============================================================
        # LANGUES + CENTRES D INTERET
        # ============================================================
        langues = data.get("langues", [])
        centres = data.get("centres_interet", [])

        if langues or centres:
            ajouter_titre_section(doc, titles["informations"])

            if langues:
                p = doc.add_paragraph()
                r = p.add_run(titles["languages_label"])
                r.bold = True
                r.font.size = Pt(9)
                p.add_run(", ".join(langues)).font.size = Pt(9)
                p.paragraph_format.space_after = Pt(2)

            if centres:
                p = doc.add_paragraph()
                r = p.add_run(titles["interests_label"])
                r.bold = True
                r.font.size = Pt(9)
                p.add_run(", ".join(centres)).font.size = Pt(9)

        logger.info(
            "word_cv_built",
            nom=nom_complet,
            nb_experiences=len(experiences),
            nb_formations=len(formations),
            country=locale.get("country", "unknown"),
        )

        return doc
