"""
Generateur Word : Brouillon de relance apres candidature.

Format email professionnel court, a envoyer manuellement par l utilisateur.
"""

from typing import Any, Dict

from docx import Document
from docx.shared import Pt

from app.core.logging import get_logger
from app.services.generation.base import DocumentGenerator
from app.services.generation.style import (
    COULEUR_PRIMAIRE,
    ajouter_paragraphe,
    configurer_marges,
    configurer_style_normal,
)

logger = get_logger("generation.word_relance")


class WordRelanceGenerator(DocumentGenerator):
    name = "word_relance"

    def build(self, data: Dict[str, Any], locale: Dict[str, Any]) -> Document:
        doc = Document()
        configurer_marges(doc, top=1.0, bottom=1.0, left=1.0, right=1.0)
        configurer_style_normal(doc, taille=11)

        # En-tete
        ajouter_paragraphe(
            doc,
            "BROUILLON DE RELANCE",
            taille=14,
            gras=True,
            couleur=COULEUR_PRIMAIRE,
            espace_apres=12,
        )

        # Metadata
        company = data.get("company_name", "")
        job = data.get("job_title", "")
        wait_days = data.get("wait_days", 7)

        ajouter_paragraphe(doc, f"Entreprise : {company}", taille=10, espace_apres=2)
        ajouter_paragraphe(doc, f"Poste : {job}", taille=10, espace_apres=2)
        ajouter_paragraphe(
            doc,
            f"Relance apres : {wait_days} jours",
            taille=10,
            italique=True,
            espace_apres=16,
        )

        # Objet
        p = doc.add_paragraph()
        r = p.add_run("Objet : ")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = COULEUR_PRIMAIRE
        objet = data.get("subject", f"Relance concernant ma candidature au poste de {job}")
        p.add_run(objet).font.size = Pt(11)
        p.paragraph_format.space_after = Pt(16)

        # Corps
        corps = data.get("body", "")
        if corps:
            for para in corps.split("\n\n"):
                if para.strip():
                    p = doc.add_paragraph()
                    r = p.add_run(para.strip())
                    r.font.size = Pt(11)
                    p.paragraph_format.space_after = Pt(10)
                    p.paragraph_format.line_spacing = 1.15

        # Signature
        signature = data.get("signature", "")
        if signature:
            ajouter_paragraphe(doc, "", taille=11)
            for ligne in signature.split("\n"):
                ajouter_paragraphe(doc, ligne, taille=11, espace_apres=0)

        logger.info(
            "word_relance_built",
            company=company,
            job=job,
            wait_days=wait_days,
        )

        return doc
