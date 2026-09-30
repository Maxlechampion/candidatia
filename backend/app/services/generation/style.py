"""Helpers de mise en forme Word (python-docx)."""

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# PALETTE DE COULEURS
# ============================================================

COULEUR_PRIMAIRE = RGBColor(27, 54, 93)      # Bleu nuit
COULEUR_SECONDAIRE = RGBColor(80, 80, 80)    # Gris anthracite
COULEUR_ACCENT = RGBColor(79, 70, 229)       # Indigo
COULEUR_NOIR = RGBColor(0, 0, 0)
COULEUR_BLANC = RGBColor(255, 255, 255)


# ============================================================
# POLICES
# ============================================================

POLICE_PRINCIPALE = "Calibri"
POLICE_ALTERNATIVE = "Arial"


# ============================================================
# HELPERS
# ============================================================


def configurer_marges(doc, top=0.6, bottom=0.6, left=0.6, right=0.6) -> None:
    """Configure les marges de toutes les sections du document."""
    for section in doc.sections:
        section.top_margin = Inches(top)
        section.bottom_margin = Inches(bottom)
        section.left_margin = Inches(left)
        section.right_margin = Inches(right)


def configurer_style_normal(doc, police=POLICE_PRINCIPALE, taille=10, couleur=COULEUR_SECONDAIRE) -> None:
    """Configure le style Normal du document."""
    style = doc.styles["Normal"]
    style.font.name = police
    style.font.size = Pt(taille)
    style.font.color.rgb = couleur


def ajouter_bordure_inferieure(paragraphe, couleur_hex="1B365D", epaisseur="6") -> None:
    """Ajoute une bordure inferieure a un paragraphe."""
    pPr = paragraphe._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), epaisseur)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), couleur_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def ajouter_paragraphe(doc, texte="", taille=10, gras=False, italique=False,
                       couleur=None, alignement=None, espace_avant=0,
                       espace_apres=0, indent_gauche=0):
    """Ajoute un paragraphe avec mise en forme."""
    p = doc.add_paragraph()

    if texte:
        run = p.add_run(texte)
        run.bold = gras
        run.italic = italique
        run.font.size = Pt(taille)
        if couleur:
            run.font.color.rgb = couleur

    if alignement is not None:
        p.alignment = alignement

    p.paragraph_format.space_before = Pt(espace_avant)
    p.paragraph_format.space_after = Pt(espace_apres)
    if indent_gauche:
        p.paragraph_format.left_indent = Inches(indent_gauche)

    return p


def ajouter_titre_section(doc, texte, couleur=COULEUR_PRIMAIRE, taille=11):
    """Ajoute un titre de section avec bordure inferieure."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(texte.upper())
    run.bold = True
    run.font.size = Pt(taille)
    run.font.color.rgb = couleur
    ajouter_bordure_inferieure(p)
    return p


def ajouter_puce(doc, texte, indent=0.15):
    """Ajoute une puce."""
    p = doc.add_paragraph(texte, style="List Bullet")
    p.paragraph_format.left_indent = Inches(indent)
    return p


def ajouter_espace(doc, points=6):
    """Ajoute un espace vertical."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("")
    run.font.size = Pt(points)
    return p


def ajouter_separateur(doc, couleur_hex="CCCCCC"):
    """Ajoute un separateur horizontal discret."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    ajouter_bordure_inferieure(p, couleur_hex=couleur_hex, epaisseur="4")
    return p


# ============================================================
# ENUM ALIGNEMENT
# ============================================================

ALIGN_GAUCHE = WD_ALIGN_PARAGRAPH.LEFT
ALIGN_CENTRE = WD_ALIGN_PARAGRAPH.CENTER
ALIGN_DROITE = WD_ALIGN_PARAGRAPH.RIGHT
ALIGN_JUSTIFIE = WD_ALIGN_PARAGRAPH.JUSTIFY
