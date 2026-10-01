"""Generateur Word : Guide de preparation a l entretien (multilingue)."""

from typing import Any, Dict

from docx import Document
from docx.shared import Pt

from app.core.logging import get_logger
from app.services.generation.base import DocumentGenerator
from app.services.generation.style import (
    ALIGN_CENTRE,
    COULEUR_PRIMAIRE,
    COULEUR_SECONDAIRE,
    ajouter_paragraphe,
    ajouter_puce,
    ajouter_titre_section,
    configurer_marges,
    configurer_style_normal,
)

logger = get_logger("generation.word_guide")


# Traductions des titres par langue
TITLES = {
    "fr": {
        "main_title": "GUIDE DE PREPARATION A L ENTRETIEN",
        "challenges": "1. Analyse des enjeux caches de l entreprise",
        "challenges_intro": "Points sur lesquels insister lors de l entretien :",
        "questions": "2. Questions probables et strategies de reponse",
        "question_label": "Question",
        "recruiter_intent": "Ce que le recruteur cherche : ",
        "your_strategy": "Votre strategie gagnante : ",
        "your_questions": "3. Vos questions pour marquer des points",
        "your_questions_intro": "Posez au moins deux de ces questions en fin d entretien :",
    },
    "en": {
        "main_title": "INTERVIEW PREPARATION GUIDE",
        "challenges": "1. Analysis of Hidden Challenges",
        "challenges_intro": "Points to emphasize during the interview:",
        "questions": "2. Likely Questions and Answer Strategies",
        "question_label": "Question",
        "recruiter_intent": "What the recruiter is looking for: ",
        "your_strategy": "Your winning strategy: ",
        "your_questions": "3. Your Questions to Score Points",
        "your_questions_intro": "Ask at least two of these questions at the end of the interview:",
    },
    "es": {
        "main_title": "GUIA DE PREPARACION PARA LA ENTREVISTA",
        "challenges": "1. Analisis de los desafios ocultos de la empresa",
        "challenges_intro": "Puntos en los que insistir durante la entrevista:",
        "questions": "2. Preguntas probables y estrategias de respuesta",
        "question_label": "Pregunta",
        "recruiter_intent": "Lo que busca el reclutador: ",
        "your_strategy": "Tu estrategia ganadora: ",
        "your_questions": "3. Tus preguntas para marcar puntos",
        "your_questions_intro": "Haz al menos dos de estas preguntas al final de la entrevista:",
    },
    "de": {
        "main_title": "LEITFADEN ZUR VORBEREITUNG AUF DAS VORSTELLUNGSGESPRAECH",
        "challenges": "1. Analyse der verborgenen Herausforderungen",
        "challenges_intro": "Punkte, die Sie im Gespraech betonen sollten:",
        "questions": "2. Wahrscheinliche Fragen und Antwortstrategien",
        "question_label": "Frage",
        "recruiter_intent": "Was der Recruiter sucht: ",
        "your_strategy": "Ihre Erfolgsstrategie: ",
        "your_questions": "3. Ihre Fragen, um Punkte zu sammeln",
        "your_questions_intro": "Stellen Sie mindestens zwei dieser Fragen am Ende des Gespraechs:",
    },
}


def _get_titles(locale: Dict[str, Any]) -> Dict[str, str]:
    """Retourne les titres selon la langue (fallback sur l anglais)."""
    language_code = "en"  # Default

    if locale:
        # Essayer plusieurs cles possibles
        for key in ("language_code", "language", "code"):
            if key in locale and locale[key]:
                language_code = str(locale[key]).lower()
                break

    titles = TITLES.get(language_code, TITLES["en"])
    logger.info("guide_titles_resolved", language_code=language_code, main_title=titles["main_title"][:30])
    return titles


class WordGuideGenerator(DocumentGenerator):
    name = "word_guide"

    def build(self, data: Dict[str, Any], locale: Dict[str, Any]) -> Document:
        titles = _get_titles(locale)

        doc = Document()
        configurer_marges(doc, top=1.0, bottom=1.0, left=1.0, right=1.0)
        configurer_style_normal(doc, taille=11)

        # ============================================================
        # TITRE PRINCIPAL
        # ============================================================
        titre_poste = data.get("titre_poste", "Poste")
        nom_entreprise = data.get("nom_entreprise", "Entreprise")

        p_titre = doc.add_paragraph()
        r_titre = p_titre.add_run(titles["main_title"])
        r_titre.bold = True
        r_titre.font.size = Pt(16)
        r_titre.font.color.rgb = COULEUR_PRIMAIRE
        p_titre.alignment = ALIGN_CENTRE
        p_titre.paragraph_format.space_after = Pt(4)

        p_sous_titre = doc.add_paragraph()
        r_st = p_sous_titre.add_run(f"{titre_poste} - {nom_entreprise}")
        r_st.italic = True
        r_st.font.size = Pt(12)
        r_st.font.color.rgb = COULEUR_SECONDAIRE
        p_sous_titre.alignment = ALIGN_CENTRE
        p_sous_titre.paragraph_format.space_after = Pt(20)

        # ============================================================
        # 1. ENJEUX CACHES
        # ============================================================
        defis = data.get("defis_cles_entreprise", [])
        if defis:
            ajouter_titre_section(doc, titles["challenges"])

            ajouter_paragraphe(
                doc,
                titles["challenges_intro"],
                taille=10,
                italique=True,
                espace_apres=6,
            )

            for defi in defis:
                ajouter_puce(doc, defi, indent=0.15)

        # ============================================================
        # 2. QUESTIONS PROBABLES
        # ============================================================
        questions = data.get("questions_probables", [])
        if questions:
            ajouter_titre_section(doc, titles["questions"])

            for i, q in enumerate(questions, 1):
                # Question
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(10)
                p_q.paragraph_format.space_after = Pt(4)
                r_q = p_q.add_run(f"{titles['question_label']} {i} : ")
                r_q.bold = True
                r_q.font.size = Pt(11)
                r_q.font.color.rgb = COULEUR_PRIMAIRE
                r_qtext = p_q.add_run(q.get("question", ""))
                r_qtext.bold = True
                r_qtext.font.size = Pt(11)

                # Intention du recruteur
                p_int = doc.add_paragraph()
                p_int.paragraph_format.left_indent = Pt(10)
                p_int.paragraph_format.space_after = Pt(3)
                r_int_l = p_int.add_run(titles["recruiter_intent"])
                r_int_l.italic = True
                r_int_l.font.size = Pt(10)
                r_int_l.font.color.rgb = COULEUR_SECONDAIRE
                r_int_v = p_int.add_run(q.get("intention_recruteur", ""))
                r_int_v.italic = True
                r_int_v.font.size = Pt(10)

                # Strategie de reponse
                p_strat = doc.add_paragraph()
                p_strat.paragraph_format.left_indent = Pt(10)
                p_strat.paragraph_format.space_after = Pt(4)
                r_strat_l = p_strat.add_run(titles["your_strategy"])
                r_strat_l.bold = True
                r_strat_l.font.size = Pt(10)
                r_strat_l.font.color.rgb = COULEUR_PRIMAIRE
                r_strat_v = p_strat.add_run(q.get("strategie_reponse", ""))
                r_strat_v.font.size = Pt(10)

        # ============================================================
        # 3. QUESTIONS A POSER
        # ============================================================
        questions_poser = data.get("questions_a_poser", [])
        if questions_poser:
            ajouter_titre_section(doc, titles["your_questions"])

            ajouter_paragraphe(
                doc,
                titles["your_questions_intro"],
                taille=10,
                italique=True,
                espace_apres=6,
            )

            for q in questions_poser:
                ajouter_puce(doc, q, indent=0.15)

        logger.info(
            "word_guide_built",
            poste=titre_poste,
            entreprise=nom_entreprise,
            nb_questions=len(questions),
        )

        return doc
