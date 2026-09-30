"""Extracteur Word .docx (python-docx + docx2txt)."""

import io

from app.core.logging import get_logger
from app.services.ingestion.base import DocumentExtractor

logger = get_logger("ingestion.docx")


class DOCXExtractor(DocumentExtractor):
    name = "docx"
    extensions = ["docx"]

    def extract(self, content_bytes: bytes, filename: str) -> str:
        text = self._extract_python_docx(content_bytes)

        if not text or len(text.strip()) < 20:
            logger.info("docx_fallback_docx2txt", filename=filename)
            text = self._extract_docx2txt(content_bytes)

        return text

    @staticmethod
    def _extract_python_docx(content_bytes: bytes) -> str:
        try:
            from docx import Document
            doc = Document(io.BytesIO(content_bytes))
            parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:
            logger.warning("python_docx_failed", error=str(e))
            return ""

    @staticmethod
    def _extract_docx2txt(content_bytes: bytes) -> str:
        try:
            import docx2txt
            return docx2txt.process(io.BytesIO(content_bytes))
        except Exception as e:
            logger.warning("docx2txt_failed", error=str(e))
            return ""
